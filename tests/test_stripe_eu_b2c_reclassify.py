"""EU B2C at Spanish 21%, stored-row reclassification, frozen declared reports,
the eur_default foreign-customer warning and the art. 73 LIVA threshold tracker.

All data is synthetic (example.* emails, round amounts).
"""
from __future__ import annotations

import json
import sqlite3
from datetime import datetime

import openpyxl
import pytest

from src.classifier import classify_payment, eur_default_foreign_warning, foreign_customer_hint
from src.database import init_db, load_classified_payments, upsert_payments
from src.declared_reports import (
    declared_vs_live_drift,
    freeze_report,
    get_declared_report,
)
from src.excel_exporter import assert_currency_geo_consistent, create_excel_report
from src.exceptions import ReportAlreadyFrozenError, StaleClassificationError
from src.models import Payment
from src.reclassify import reclassify_stored
from src.tax_engine import (
    compute_eu_b2c_threshold,
    compute_modelo_130,
    compute_modelo_303,
    compute_modelo_349,
    compute_oss_return,
)
from src.vat_rules import vat_amount_on_base, vat_base_from_inclusive, vat_treatment


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

@pytest.fixture
def db_path(tmp_path):
    path = tmp_path / "accounting.db"
    init_db(path)
    return path


@pytest.fixture
def conn(db_path):
    c = sqlite3.connect(str(db_path))
    c.row_factory = sqlite3.Row
    yield c
    c.close()


def _payment(pid: str, when: str, amount: float, desc: str = "Subscription creation",
             currency: str = "eur", **kw) -> Payment:
    return Payment(
        id=pid, created_date=when, converted_amount=amount, converted_amount_refunded=0.0,
        description=desc, fee=0.0, currency=currency, **kw,
    )


def _store(db_path, payments: list[Payment], rules: dict) -> None:
    """Persist payments and their current classification."""
    from src.database import upsert_classified

    upsert_payments(payments, db_path=db_path)
    upsert_classified([classify_payment(p, rules) for p in payments], db_path=db_path)


def _set_stale(conn, pid: str, geo: str, geo_rule: str) -> None:
    conn.execute(
        "UPDATE transactions SET geo_region = ?, geo_rule = ? WHERE id = ?", (geo, geo_rule, pid)
    )
    conn.commit()


# ---------------------------------------------------------------------------
# EU B2C → Spanish 21% (not OSS-registered)
# ---------------------------------------------------------------------------

class TestEuB2CSpanish21:
    def test_default_eu_newsletter_is_eu_b2c_es21(self):
        assert vat_treatment("NEWSLETTER", "EU_NOT_SPAIN") == "EU_B2C_ES21"
        assert vat_treatment("NEWSLETTER", "EU_NOT_SPAIN", {"tax": {}}) == "EU_B2C_ES21"

    def test_oss_only_when_registered(self):
        cfg = {"tax": {"oss_registered": True}}
        assert vat_treatment("NEWSLETTER", "EU_NOT_SPAIN", cfg) == "OSS_EU"

    def test_oss_setting_without_registration_is_coerced(self):
        cfg = {"tax": {"oss_registered": False, "default_vat_treatment_eu_newsletter": "OSS_EU"}}
        assert vat_treatment("NEWSLETTER", "EU_NOT_SPAIN", cfg) == "EU_B2C_ES21"

    def test_not_vat_registered_is_exempt(self):
        cfg = {"tax": {"vat_registered": False}}
        assert vat_treatment("NEWSLETTER", "EU_NOT_SPAIN", cfg) == "IVA_EXEMPT"

    def test_eu_coaching_and_illustrations_without_vat_id_default_to_b2c(self):
        # #113: EU sales follow the customer's status, not the activity — with
        # no VAT id on file, coaching/illustrations default to B2C (21%) just
        # like newsletter, instead of the old always-B2B default.
        assert vat_treatment("COACHING", "EU_NOT_SPAIN") == "EU_B2C_ES21"
        assert vat_treatment("ILLUSTRATIONS", "EU_NOT_SPAIN") == "EU_B2C_ES21"

    def test_eu_coaching_and_illustrations_with_vat_id_is_b2b(self):
        assert vat_treatment("COACHING", "EU_NOT_SPAIN", buyer_vat_id="DE123456789") == "IVA_EU_B2B"
        assert vat_treatment("ILLUSTRATIONS", "EU_NOT_SPAIN", buyer_vat_id="FR12345678901") == "IVA_EU_B2B"
        # Blank/whitespace-only VAT id is treated as unknown.
        assert vat_treatment("COACHING", "EU_NOT_SPAIN", buyer_vat_id="  ") == "EU_B2C_ES21"

    def test_legacy_b2b_config_default_no_longer_forces_b2b(self):
        # A stale config value from before #113 is accepted (doesn't crash) but
        # ignored: B2B can no longer be forced without a known customer VAT id.
        cfg = {"tax": {"default_vat_treatment_eu_coaching": "IVA_EU_B2B"}}
        assert vat_treatment("COACHING", "EU_NOT_SPAIN", cfg) == "EU_B2C_ES21"

    def test_base_and_cuota_at_21(self):
        base = vat_base_from_inclusive(121.0, "EU_B2C_ES21")
        assert base == pytest.approx(100.0)
        assert vat_amount_on_base(base, "EU_B2C_ES21") == pytest.approx(21.0)

    def test_modelo_303_and_130_book_eu_b2c_at_21(self, db_path, conn, sample_rules):
        # An EUR newsletter charge from an EU consumer → EU_NOT_SPAIN via eur_newsletter_default.
        _store(db_path, [_payment("ch_eu1", "2025-02-10T10:00:00", 121.0,
                                  email_meta="reader@example.at")], sample_rules)
        r303 = compute_modelo_303(2025, 1, conn, {"tax": {}})
        assert r303.box_01_base == pytest.approx(100.0)
        assert r303.box_03_cuota == pytest.approx(21.0)
        assert r303.oss_base == 0.0
        audit_01 = next(a for a in r303.audit if a.cell == "c07_base")
        recs = json.loads(audit_01.inputs_json)["records"]
        assert any(r.get("vat_treatment") == "EU_B2C_ES21" for r in recs)

        r130 = compute_modelo_130(2025, 1, conn, {"tax": {}})
        assert r130.box_01_ingresos == pytest.approx(100.0)

        oss = compute_oss_return(2025, 1, conn, {"tax": {}})
        assert oss.rows == []
        assert any(a.cell == "oss_not_registered" for a in oss.audit)

    def test_eu_coaching_without_vat_id_books_at_21_not_b2b(self, db_path, conn, sample_rules):
        # Acceptance criterion (accounting-quarterly#113): an EU coaching sale
        # with no customer VAT id on file is 21% Spanish IVA, not B2B.
        _store(db_path, [_payment("ch_eu_coach", "2025-02-10T10:00:00", 121.0,
                                  desc="Calendly coaching", email_meta="test@example.de")],
               sample_rules)
        r303 = compute_modelo_303(2025, 1, conn, {"tax": {}})
        assert r303.box_01_base == pytest.approx(100.0)
        assert r303.box_03_cuota == pytest.approx(21.0)
        assert r303.box_59_intracom_entregas == 0.0
        assert compute_modelo_349(2025, 1, conn, {"tax": {}}).rows == []

    def test_eu_coaching_with_vat_id_is_b2b_and_lists_on_349(self, db_path, conn, sample_rules):
        # Acceptance criterion (accounting-quarterly#113): an EU coaching sale
        # with a customer VAT id on file is EU B2B, listed on Modelo 349 key S.
        rules = json.loads(json.dumps(sample_rules))
        rules["customer_vat_ids"] = {
            "email_vat_ids": {"test@example.de": "DE123456789"}, "name_vat_ids": {},
        }
        _store(db_path, [_payment("ch_eu_coach_b2b", "2025-02-10T10:00:00", 500.0,
                                  desc="Calendly coaching", email_meta="test@example.de")],
               rules)

        r303 = compute_modelo_303(2025, 1, conn, {"tax": {}})
        assert r303.box_01_base == 0.0
        assert r303.box_59_intracom_entregas == pytest.approx(500.0)

        r349 = compute_modelo_349(2025, 1, conn, {"tax": {}})
        assert len(r349.rows) == 1
        assert r349.rows[0].vat_id == "DE123456789"
        assert r349.rows[0].base == pytest.approx(500.0)


# ---------------------------------------------------------------------------
# Non-EUR charges can never be exported as EUR / EU_NOT_SPAIN
# ---------------------------------------------------------------------------

class TestNonEurNeverExportedAsEu:
    @pytest.mark.parametrize("currency", ["usd", "aud"])
    def test_classifier_routes_non_eur_newsletter_outside_eu(self, sample_rules, currency):
        cp = classify_payment(_payment("ch_x", "2026-01-16T10:00:00", 10.0, currency=currency),
                              sample_rules)
        assert cp.geo_region == "OUTSIDE_EU"
        assert cp.geo_rule == f"non_eur_currency:{currency}"

    @pytest.mark.parametrize("currency", ["usd", "aud"])
    def test_stale_row_blocks_export_until_reclassified(
        self, db_path, conn, sample_rules, tmp_path, currency
    ):
        _store(db_path, [_payment("ch_fx", "2026-01-29T10:00:00", 60.0, currency=currency)],
               sample_rules)
        # Simulate a row classified back when its currency was still recorded as EUR.
        _set_stale(conn, "ch_fx", "EU_NOT_SPAIN", "eur_newsletter_default")

        stale = load_classified_payments(db_path=db_path)
        with pytest.raises(StaleClassificationError):
            create_excel_report(stale, tmp_path / "r.xlsx", 2026, 1)

        reclassify_stored(datetime(2026, 1, 1), rules=sample_rules, db_path=db_path)
        fresh = load_classified_payments(db_path=db_path)
        out = create_excel_report(fresh, tmp_path / "r.xlsx", 2026, 1)

        ws = openpyxl.load_workbook(out)["import"]
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        assert row["Currency"] == currency.upper()
        assert row["Geo Region"] == "OUTSIDE_EU"


# ---------------------------------------------------------------------------
# EU non-euro currencies classified by charge country (#111)
# ---------------------------------------------------------------------------

class TestEuNonEuroCurrencyByCountry:
    @pytest.mark.parametrize("currency", ["dkk", "sek", "pln"])
    def test_eu_non_euro_newsletter_with_eu_card_is_eu_b2c_es21(self, sample_rules, currency):
        # A DKK/SEK/PLN newsletter charge from an EU card is an EU B2C sale,
        # not outside the EU — Spanish 21% (art. 73 LIVA, not OSS-registered).
        cp = classify_payment(
            _payment("ch_eu_nc", "2026-02-10T10:00:00", 121.0, currency=currency, card_country="DE"),
            sample_rules,
        )
        assert cp.geo_region == "EU_NOT_SPAIN"
        assert cp.geo_rule == "country:DE"
        treatment = vat_treatment(cp.activity_type, cp.geo_region)
        assert treatment == "EU_B2C_ES21"
        base = vat_base_from_inclusive(cp.net_amount, treatment)
        assert base == pytest.approx(100.0)
        assert vat_amount_on_base(base, treatment) == pytest.approx(21.0)

    def test_usd_with_us_card_still_outside_eu(self, sample_rules):
        cp = classify_payment(
            _payment("ch_us_card", "2026-02-11T10:00:00", 60.0, currency="usd", card_country="US"),
            sample_rules,
        )
        assert cp.geo_region == "OUTSIDE_EU"
        assert cp.geo_rule == "country:US"
        assert vat_treatment(cp.activity_type, cp.geo_region) == "IVA_EXPORT"

    def test_export_accepts_country_rule_on_non_eur_row(self, sample_rules):
        # A non-EUR row classified via the new charge-country rule is a
        # legitimate, currency-agnostic classification — not a stale one.
        cp = classify_payment(
            _payment("ch_dkk_eu", "2026-02-12T10:00:00", 100.0, currency="dkk", card_country="DK"),
            sample_rules,
        )
        assert cp.geo_rule == "country:DK"
        assert_currency_geo_consistent([cp])  # must not raise

    def test_export_still_rejects_eur_default_on_non_eur_row(self, sample_rules):
        # #94's original protection still holds: a genuinely stale EUR-branch
        # rule on a non-EUR row is refused.
        cp = classify_payment(_payment("ch_stale", "2026-02-13T10:00:00", 100.0), sample_rules)
        stale = cp.model_copy(update={"currency": "usd", "geo_rule": "eur_default"})
        with pytest.raises(StaleClassificationError):
            assert_currency_geo_consistent([stale])


# ---------------------------------------------------------------------------
# Reclassify stored rows
# ---------------------------------------------------------------------------

class TestReclassify:
    def test_changes_stale_row_idempotently(self, db_path, conn, sample_rules):
        _store(db_path, [
            _payment("ch_ok", "2025-04-02T10:00:00", 50.0, desc="Calendly coaching"),
            _payment("ch_old", "2025-05-02T10:00:00", 80.0, desc="Charge for John Doe portrait"),
        ], sample_rules)
        # Rows stored before the "john doe" override existed.
        _set_stale(conn, "ch_old", "SPAIN", "eur_default")

        dry = reclassify_stored(datetime(2025, 1, 1), dry_run=True, rules=sample_rules,
                                db_path=db_path)
        assert [c.id for c in dry.changes] == ["ch_old"]
        assert dry.changes[0].old[1] == "SPAIN"
        assert dry.changes[0].new[1] == "OUTSIDE_EU"
        assert dry.by_quarter() == {(2025, 2): 1}
        # Dry run wrote nothing.
        assert conn.execute("SELECT geo_region FROM transactions WHERE id='ch_old'").fetchone()[0] \
            == "SPAIN"

        applied = reclassify_stored(datetime(2025, 1, 1), rules=sample_rules, db_path=db_path)
        assert len(applied.changes) == 1
        row = conn.execute("SELECT geo_region, geo_rule FROM transactions WHERE id='ch_old'").fetchone()
        assert tuple(row) == ("OUTSIDE_EU", "name_override:john doe")

        again = reclassify_stored(datetime(2025, 1, 1), rules=sample_rules, db_path=db_path)
        assert again.changes == []
        assert again.scanned == 2

    def test_respects_from_date(self, db_path, conn, sample_rules):
        _store(db_path, [_payment("ch_early", "2024-03-01T10:00:00", 10.0, currency="usd")],
               sample_rules)
        _set_stale(conn, "ch_early", "EU_NOT_SPAIN", "eur_newsletter_default")
        result = reclassify_stored(datetime(2024, 6, 1), rules=sample_rules, db_path=db_path)
        assert result.scanned == 0 and result.changes == []


# ---------------------------------------------------------------------------
# Frozen declared report
# ---------------------------------------------------------------------------

class TestDeclaredReport:
    def _freeze_q1(self, db_path, conn, sample_rules, tmp_path, supersede=False):
        payments = load_classified_payments(datetime(2026, 1, 1), datetime(2026, 3, 31, 23, 59, 59),
                                            db_path=db_path)
        path = create_excel_report(payments, tmp_path / "Stripe_Report_Q1_2026.xlsx", 2026, 1)
        return freeze_report(conn, 2026, 1, payments, path, supersede=supersede)

    def test_engine_uses_frozen_amounts_after_fx_change(self, db_path, conn, sample_rules, tmp_path):
        _store(db_path, [
            _payment("ch_es", "2026-02-01T10:00:00", 121.0, desc="Calendly coaching"),
            _payment("ch_usd", "2026-02-15T10:00:00", 60.0, currency="usd",
                     amount_original=70.0, fx_rate=1.1667),
        ], sample_rules)
        report = self._freeze_q1(db_path, conn, sample_rules, tmp_path)
        assert report.version == 1 and report.n_transactions == 2
        assert report.total_net_eur == pytest.approx(181.0)
        assert len(report.sha256) == 64

        before_130 = compute_modelo_130(2026, 1, conn, {"tax": {}}).box_01_ingresos
        assert before_130 == pytest.approx(100.0 + 60.0)

        # A later re-fetch re-converts the USD charge at a different ECB rate.
        conn.execute("UPDATE transactions SET converted_amount = 61.37 WHERE id = 'ch_usd'")
        conn.execute("UPDATE transactions SET converted_amount = 150.0 WHERE id = 'ch_es'")
        conn.commit()

        r303 = compute_modelo_303(2026, 1, conn, {"tax": {}})
        assert r303.box_01_base == pytest.approx(100.0)
        assert r303.export_base == pytest.approx(60.0)
        assert compute_modelo_130(2026, 1, conn, {"tax": {}}).box_01_ingresos == \
            pytest.approx(before_130)

        live = load_classified_payments(datetime(2026, 1, 1), datetime(2026, 3, 31, 23, 59, 59),
                                        db_path=db_path)
        drift = declared_vs_live_drift(conn, 2026, 1, live)
        assert drift["amount_differs"] == ["ch_es", "ch_usd"]

    def test_declared_rows_are_immutable(self, db_path, conn, sample_rules, tmp_path):
        _store(db_path, [_payment("ch_1", "2026-01-10T10:00:00", 10.0)], sample_rules)
        self._freeze_q1(db_path, conn, sample_rules, tmp_path)
        with pytest.raises(sqlite3.DatabaseError):
            conn.execute("UPDATE declared_report_lines SET converted_amount = 1.0")
        with pytest.raises(sqlite3.DatabaseError):
            conn.execute("DELETE FROM declared_reports")

    def test_refreeze_requires_supersede(self, db_path, conn, sample_rules, tmp_path):
        _store(db_path, [_payment("ch_1", "2026-01-10T10:00:00", 10.0)], sample_rules)
        self._freeze_q1(db_path, conn, sample_rules, tmp_path)
        with pytest.raises(ReportAlreadyFrozenError):
            self._freeze_q1(db_path, conn, sample_rules, tmp_path)

        conn.execute("UPDATE transactions SET converted_amount = 12.0 WHERE id = 'ch_1'")
        conn.commit()
        v2 = self._freeze_q1(db_path, conn, sample_rules, tmp_path, supersede=True)
        assert v2.version == 2
        assert get_declared_report(conn, 2026, 1).version == 2
        # The engine follows the latest declared version.
        assert compute_modelo_130(2026, 1, conn, {"tax": {}}).box_01_ingresos == \
            pytest.approx(round(12.0 / 1.21, 2))

    def test_no_declared_report_leaves_live_amounts(self, db_path, conn, sample_rules):
        _store(db_path, [_payment("ch_1", "2026-01-10T10:00:00", 121.0, desc="Calendly coaching")],
               sample_rules)
        assert get_declared_report(conn, 2026, 1) is None
        assert compute_modelo_303(2026, 1, conn, {"tax": {}}).box_01_base == pytest.approx(100.0)


# ---------------------------------------------------------------------------
# eur_default foreign-customer warning
# ---------------------------------------------------------------------------

class TestForeignCustomerWarning:
    def _coaching(self, **kw) -> Payment:
        return _payment("ch_w", "2026-01-10T10:00:00", 400.0, desc="Calendly coaching", **kw)

    def test_foreign_email_domain_on_eur_default_warns(self, sample_rules):
        cp = classify_payment(self._coaching(email_meta="a@example.nl"), sample_rules)
        assert cp.geo_rule == "eur_default" and cp.geo_region == "SPAIN"
        warning = eur_default_foreign_warning(cp)
        assert warning and ".nl" in warning

    def test_email_in_description_and_card_country(self, sample_rules):
        p = _payment("ch_w", "2026-01-10T10:00:00", 400.0,
                     desc="Charge for a@example.nl artwork", card_country="NL")
        hint = foreign_customer_hint(p)
        assert "card country NL" in hint and "email domain .nl" in hint

    def test_billing_country_from_raw_source(self):
        p = self._coaching(raw_source={"billing_details": {"address": {"country": "DE"}}})
        assert foreign_customer_hint(p) == "billing country DE"

    @pytest.mark.parametrize("email", ["a@example.es", "a@example.com", "a@example.io", None])
    def test_spanish_or_generic_domain_does_not_warn(self, sample_rules, email):
        cp = classify_payment(self._coaching(email_meta=email, card_country="ES"), sample_rules)
        assert eur_default_foreign_warning(cp) is None

    def test_only_eur_default_rule_warns(self, sample_rules):
        # Override-matched row: the user already decided, no warning.
        cp = classify_payment(self._coaching(email_meta="test@example.de"), sample_rules)
        assert cp.geo_rule.startswith("email_override")
        assert eur_default_foreign_warning(cp) is None


# ---------------------------------------------------------------------------
# EU B2C threshold tracker (art. 73 LIVA)
# ---------------------------------------------------------------------------

class TestEuB2CThreshold:
    def _newsletters(self, db_path, sample_rules, year: int, gross_each: float, n: int) -> None:
        _store(db_path, [
            _payment(f"ch_{year}_{i}", f"{year}-0{1 + i % 3}-1{i % 9}T10:00:00", gross_each,
                     email_meta=f"reader{i}@example.at")
            for i in range(n)
        ], sample_rules)

    def test_ok_below_warning(self, db_path, conn, sample_rules):
        self._newsletters(db_path, sample_rules, 2026, 121.0, 10)  # base 1,000
        t = compute_eu_b2c_threshold(2026, 1, conn, {"tax": {}})
        assert t.ytd_base_eur == pytest.approx(1000.0)
        assert t.n_transactions == 10
        assert t.status == "OK"

    def test_warns_at_80_percent(self, db_path, conn, sample_rules):
        self._newsletters(db_path, sample_rules, 2026, 1210.0, 7)  # base 7,000
        _store(db_path, [_payment("ch_big", "2026-03-20T10:00:00", 1210.0,
                                  email_meta="reader@example.fr")], sample_rules)  # → 8,000
        t = compute_eu_b2c_threshold(2026, 1, conn, {"tax": {}})
        assert t.ytd_base_eur == pytest.approx(8000.0)
        assert t.ratio == pytest.approx(0.8)
        assert t.status == "WARNING"

    def test_exceeded_this_year_or_previous_year(self, db_path, conn, sample_rules):
        self._newsletters(db_path, sample_rules, 2025, 1210.0, 9)  # 2025 base 9,000
        assert compute_eu_b2c_threshold(2026, 1, conn, {"tax": {}}).status == "OK"
        _store(db_path, [_payment("ch_2025_x", "2025-11-02T10:00:00", 2420.0,
                                  email_meta="r@example.it")], sample_rules)  # 2025 → 11,000
        t = compute_eu_b2c_threshold(2026, 1, conn, {"tax": {}})
        assert t.previous_year_base_eur == pytest.approx(11000.0)
        assert t.status == "EXCEEDED"

    def test_spain_and_non_eu_sales_do_not_count(self, db_path, conn, sample_rules):
        _store(db_path, [
            _payment("ch_es", "2026-01-05T10:00:00", 5000.0, desc="Calendly coaching"),
            _payment("ch_us", "2026-01-06T10:00:00", 9000.0, currency="usd"),
        ], sample_rules)
        t = compute_eu_b2c_threshold(2026, 1, conn, {"tax": {}})
        assert t.ytd_base_eur == 0.0 and t.status == "OK"
